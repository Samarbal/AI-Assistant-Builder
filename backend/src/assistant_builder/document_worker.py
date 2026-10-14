"""Bounded document parser subprocess; no application configuration or networking."""
import io
import json
import logging
import sys
import zipfile

MAX_TEXT = 2_000_000


def extract(content, file_format):
    if not 0 < len(content) <= 10_000_000:
        raise ValueError()
    if file_format == 'pdf':
        from pypdf import PdfReader
        if not content.startswith(b'%PDF-'):
            raise ValueError()
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError()
        parts = []
        length = 0
        for page in reader.pages:
            # Avoid unbounded decompressed content-stream parsing.
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > 5_000_000:
                raise ValueError()
            value = page.extract_text() or ''
            length += len(value)
            if length > MAX_TEXT:
                raise ValueError()
            parts.append(value)
        result = '\n'.join(parts)
    elif file_format == 'docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 50_000_000:
                raise ValueError()
            if any(e.flag_bits & 1 or e.filename.lower().endswith('vbaproject.bin') for e in entries):
                raise ValueError()
            if 'word/document.xml' not in archive.namelist():
                raise ValueError()
        document = Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        result = '\n'.join(parts)
    else:
        raise ValueError()
    result = result.replace('\x00', '').strip()
    if not result or len(result) > MAX_TEXT:
        raise ValueError()
    return result


if __name__ == '__main__':
    logging.disable(logging.CRITICAL)
    try:
        result = extract(sys.stdin.buffer.read(10_000_001), sys.argv[1])
        sys.stdout.write(json.dumps({'text': result}))
    except Exception:
        raise SystemExit(2) from None
