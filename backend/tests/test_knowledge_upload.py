from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from assistant_builder.auth import OwnerContext, get_current_owner
from assistant_builder.extraction import extract_text, InvalidDocument
from assistant_builder.knowledge_api import service
from assistant_builder.knowledge_repository import KnowledgeNotFound
from assistant_builder.knowledge_service import KnowledgeService
from assistant_builder.main import create_app
from assistant_builder.models import Base, Assistant, KnowledgeFile, KnowledgeFileText
from assistant_builder.storage import SupabaseFileStorage, StorageUnavailable


def docx_bytes():
    document = Document()
    document.add_paragraph('Our opening hours are nine to five.')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Weekend closed.'
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(encrypted=False, empty=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if not empty:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 200 Td (Our opening hours are nine to five.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted:
        writer.encrypt('test-only')
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


@pytest.mark.parametrize('kind,content', [('docx', docx_bytes()), ('pdf', pdf_bytes())], ids=['docx', 'pdf'])
def test_real_extraction_subprocess(kind, content):
    assert 'nine to five' in extract_text(content, kind)


@pytest.mark.parametrize('kind,content', [
    ('pdf', b'not a PDF'), ('pdf', pdf_bytes(encrypted=True)),
    ('pdf', pdf_bytes(empty=True)), ('docx', b'not a zip'), ('pdf', docx_bytes()),
], ids=['malformed-pdf', 'encrypted-pdf', 'empty-pdf', 'malformed-docx', 'wrong-format'])
def test_invalid_documents_fail_closed(kind, content):
    with pytest.raises(InvalidDocument):
        extract_text(content, kind)


def test_storage_does_not_upsert_or_leak_provider_errors():
    def handler(request):
        assert request.headers['x-upsert'] == 'false'
        assert request.headers['content-type'] == 'application/pdf'
        return httpx.Response(403, text='sensitive provider error')
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        storage = SupabaseFileStorage('https://example.supabase.co', 'test', 'knowledge-files', client)
        with pytest.raises(StorageUnavailable) as error:
            storage.upload('owners/test/original.pdf', b'pdf', 'pdf')
        assert 'sensitive' not in str(error.value)


@pytest.fixture
def workflow(monkeypatch):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, expire_on_commit=False)
    owner, assistant = uuid4(), uuid4()
    with sessions.begin() as session:
        session.add(Assistant(id=assistant, owner_id=owner, name='test', role='support', colors={}, response_style='short', tone='friendly', fallback_message='Unknown'))
    monkeypatch.setattr('assistant_builder.knowledge_service.set_owner_identity', lambda session, owner: None)
    calls = []
    storage = SimpleNamespace(upload=lambda *args: calls.append(args))
    yield KnowledgeService(sessions, storage), owner, assistant, sessions, calls
    engine.dispose()


def test_workflow_commits_text_and_metadata(workflow):
    flow, owner, assistant, sessions, calls = workflow
    item = flow.upload(owner, assistant, 'hours.docx', docx_bytes())
    assert item.processing_state == 'extracted' and len(calls) == 1
    with sessions() as session:
        assert 'nine to five' in session.scalar(select(KnowledgeFileText.extracted_text))
        assert session.get(KnowledgeFile, item.id).processing_state == 'extracted'
    assert str(owner) in calls[0][0] and str(assistant) in calls[0][0]


def test_invalid_file_recorded_and_never_uploaded(workflow):
    flow, owner, assistant, sessions, calls = workflow
    with pytest.raises(InvalidDocument):
        flow.upload(owner, assistant, 'bad.pdf', b'bad')
    with sessions() as session:
        item = session.scalar(select(KnowledgeFile))
        assert item.processing_state == 'failed' and item.cleanup_state == 'pending'
        assert session.scalar(select(KnowledgeFileText.file_id)) is None
    assert not calls
    flow.remove(owner, assistant, item.id)
    assert flow.list(owner, assistant) == []


def test_cross_owner_has_no_storage_side_effect(workflow):
    flow, _, assistant, _, calls = workflow
    with pytest.raises(KnowledgeNotFound):
        flow.upload(uuid4(), assistant, 'hours.docx', docx_bytes())
    assert not calls


def test_storage_failure_persists_failed_state(workflow):
    flow, owner, assistant, sessions, _ = workflow
    def fail(*args):
        raise StorageUnavailable()
    flow.storage.upload = fail
    with pytest.raises(StorageUnavailable):
        flow.upload(owner, assistant, 'hours.docx', docx_bytes())
    with sessions() as session:
        item = session.scalar(select(KnowledgeFile))
        assert item.error_code == 'storage_unavailable' and item.cleanup_state == 'pending'


def test_swagger_upload_auth_validation_and_no_owner_override():
    app = create_app()
    assistant, owner = uuid4(), uuid4()
    calls = []
    def upload(*args):
        calls.append(args)
        return SimpleNamespace(id=uuid4(), assistant_id=assistant, original_filename='file.pdf', file_format='pdf', size_bytes=4, processing_state='extracted', error_code=None, created_at=datetime.now(timezone.utc))
    app.dependency_overrides[service] = lambda: SimpleNamespace(storage=True, upload=upload)
    with TestClient(app) as client:
        path = f'/assistants/{assistant}/knowledge-files'
        assert client.post(path, files={'file': ('file.pdf', b'data')}).status_code == 401
        app.dependency_overrides[get_current_owner] = lambda: OwnerContext(owner)
        assert client.post(path, files={'file': ('file.pdf', b'data')}).status_code == 201
        assert calls[0][0] == owner
        assert client.post(path, files={'file': ('file.txt', b'data')}).status_code == 422
        assert client.post(path, files={'file': ('file.pdf', b'data')}, data={'owner_id': str(uuid4())}).status_code == 400
        assert client.post(path, content=b'', headers={'content-length': '10065537'}).status_code == 413
        schema = client.get('/openapi.json').json()
        assert 'multipart/form-data' in schema['paths']['/assistants/{assistant_id}/knowledge-files']['post']['requestBody']['content']
        assert schema['components']['securitySchemes']['HTTPBearer']['scheme'] == 'bearer'
        assert schema['paths']['/assistants/{assistant_id}/knowledge-files']['post']['security'] == [{'HTTPBearer': []}]
        assert client.get('/docs').status_code == 200
